from pathlib import Path

import xlsxwriter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent
PPT = OUT / "Ascend910B_HG-KDA汇报-20261028.pptx"
XLSX = OUT / "算子结果总表-20261028.xlsx"

NAVY = "14253D"
BLUE = "2864A5"
TEAL = "168C8C"
PALE = "EAF1F7"
PALE_TEAL = "E6F4F2"
GOLD = "D48A24"
RED = "B44949"
GRAY = "526174"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"


meeting = [
    ["value_selecting_reduction_backward", "0.4935×", "0.759×", "1.635×", "96项完整评测；扩展语义通过", "融合补零与梯度回填；修复非连续stride", "默认条件达标", "operators/会议10算子最终验收-20261026.md；campaigns/night-20261021-v3/交付清单.md"],
    ["scatter_add", "精度失败", "0.258×；扩展失败", "1.900×", "18/18正式；18/18语义；9/9性能输入；18/18动态范围", "排序、片上二分、显式分段扫描，避免跨目标大数吞掉小数", "默认条件达标；替代旧1.306×无效候选", "operators/会议剩余五算子交付-20261024.md；campaigns/scatter-segmented-20261023/README.md"],
    ["uniform", "精度通过，无有效加速比", "0.141×；扩展失败", "1.325×", "18/18正式；标准Philox逐位及随机状态检查通过", "修复无符号移位与counter进位，优化状态读取和启动", "默认条件达标；旧1.497×候选因Philox不正确被替换", "operators/会议10算子最终验收-20261026.md；campaigns/endurance-20261022-uniform-shift-repair-v2/精简汇报.md"],
    ["rad2deg_", "精度失败", "正确性失败，33/36", "1.117×", "36/36正式；18项补充通过", "保持原地语义；分块与低开销启动", "默认条件达标", "operators/会议10算子最终验收-20261026.md；operators/夜间连续优化交付-20261023.md"],
    ["upsample_bilinear2d", "精度失败", "编译／布局失败", "0.915×", "24/24正式；48＋36项补充通过", "连续读取相邻输入行；一次计算两行输出；修复分块交界重复写", "默认条件达标；实际延迟较同批父版改善2.559×", "operators/两算子收尾交付-20261024.md；campaigns/two-closure-20261024/README.md"],
    ["log10_", "精度失败", "正确性失败，33/36", "0.875×", "36/36正式；42＋12＋9项补充通过", "修复高维stride；去除中间FP32转换并优化地址键启动缓存", "默认条件达标；buffer批次同批父版0.751×，实际加速1.125×", "campaigns/buffer-pointwise-20261022/精简汇报.md；campaigns/paired-pointwise-20261022/精简汇报.md"],
    ["selu_", "精度失败", "历史固定版0.119×", "0.837×（10项验收表）；另批同合同测试0.898×", "36/36正式；42＋12＋9项补充通过", "调整分块、精度安排和启动；保留通过全部门禁的源码", "默认条件达标；两次不同批次测值分列，不合并", "operators/会议10算子最终验收-20261026.md；campaigns/buffer-pointwise-20261022/精简汇报.md"],
    ["cumulative_trapezoid", "精度失败", "历史固定版0.363×；默认扩展数值失败", "0.812×（仅官方确定性条件）", "确定性条件33/33；补充18＋9项计时输入检查", "Triton融合相邻加法／缩放，再调用CANN aclnnCumsum", "条件达标；默认参考重复结果不自洽，不能列默认达标", "operators/两算子收尾交付-20261024.md；operators/会议10算子最终验收-20261026.md"],
    ["fractional_max_pool3d", "0.0266×", "9/9正式；扩展16/27；PyTorch参考含CPU回退", "17.683×（CPU回退参考）；NPU分解对照10.066×", "两组各9/9；27项输入／布局、30项特殊语义", "连续128个输出分块；保留低精度窗口坐标中间舍入", "仅补充NPU分解对照；17.683×不能作原生NPU验收成绩", "operators/会议10算子最终验收-20261026.md；campaigns/acceptance-three-20261026/精简汇报.md"],
    ["_conj_physical", "精度失败", "新接口正式／扩展失败；PyTorch参考含CPU回退", "47.170×（CPU回退参考）；NPU分解对照2.296×", "两组各14/14；34项补充语义通过", "复用已验证实现；新DMA候选未优于原实现", "仅补充NPU分解对照；不可与conj_physical旧接口混用", "operators/会议10算子最终验收-20261026.md；campaigns/acceptance-three-20261026/精简汇报.md"],
]

sglang = [
    ["gelu_and_mul", "0.068×；旧合同19/19，扩展FP32失败", "2.271×", "正式19/19；补充18/18；性能输入15/15；缓存／偏移／stream 9/9", "匹配本机GELU公式；连续二维分块，简化索引并优化启动", "已交付；当前唯一建议继续优化的SGLang对象", "operators/夜间连续优化交付-20261023.md"],
    ["moe_sum_reduce", "0.965×", "1.624×（两次1.628×／1.619×几何平均）", "两次8项正式；36项扩展语义通过", "合并token归约；连续整除快路径，其他布局保留通用stride处理", "已交付成果；HG-KDA生成后由Codex修复布局", "campaigns/night-20261021-v3/精简汇报.md；campaigns/night-20261021-v3-moe-hg-layout-repair/README.md"],
    ["per_token_quant_int8", "2.085×", "—", "原版19/19", "固定原版体检；本阶段没有另做HG-KDA优化版", "归档体检项", "campaigns/night-20261021-v3/精简汇报.md"],
    ["rotary_embedding", "2.130×", "—", "原版9/9", "固定原版体检；本阶段没有另做HG-KDA优化版", "归档体检项", "campaigns/night-20261021-v3/精简汇报.md"],
    ["silu_and_mul_masked", "28.947×（组合参考含expert循环和.item()同步）", "—", "固定版本体检；对照口径待核", "保留为组合参考口径观察项", "不作为公平融合kernel加速结论", "campaigns/night-20261021-v3/精简汇报.md"],
]

historic = [
    ["mse_loss_backward", "1.1688×", "1.11×（原版）", "1.34×", "1.20×", "完整测试通过", "扩大分块；连续输入快路径", "operators/研究院汇报-20261014.md"],
    ["t_copy", "3.0894×", "1.96×（原版）", "3.43×", "1.73×", "9/9", "按形状选tile；多块合并处理", "operators/研究院汇报-20261014.md"],
    ["prelu_kernel_backward", "1.8152×", "原版3/21；修复版0.49×", "0.95×", "1.97×（修复版）", "完整测试通过", "融合梯度与归约；FP32误差补偿", "operators/研究院汇报-20261014.md"],
    ["batch_norm_backward", "4.3885×", "原版9/30；修复版2.32×", "5.01×", "2.19×（修复版）", "完整测试通过", "按形状分流、融合与分阶段归约", "operators/研究院汇报-20261014.md"],
    ["smooth_l1_loss_backward", "2.2139×", "原版290/291；修复版2.17×", "3.48×", "1.61×（修复版）", "完整测试通过", "修复beta=0边界；扩大分块并提前确定常量分支", "operators/研究院汇报-20261014.md"],
    ["logaddexp2", "0.3218×", "0.231×", "1.262×", "—", "36/36", "稳定公式、连续快路径、精度安排和多块处理", "operators/四算子闭环-20261014-four-refined.md"],
    ["amin", "0.5536×", "0.308×", "0.934×", "—", "27/27", "直接归约；整除时去边界判断；精度与缓冲优化", "operators/四算子闭环-20261014-four-refined.md"],
    ["matmul_bias_activation", "0.5408×", "0.439×", "0.909×", "—", "42/42", "双输出块复用输入；扩大K分块；融合bias／激活", "operators/四算子闭环-20261014-four-refined.md"],
    ["narrow_copy", "0.0138×", "约0.010×", "1.118×", "—", "33/33", "小输入CANN DMA；大输入Triton分块复制", "operators/四算子闭环-20261014-four-refined.md"],
]


def build_xlsx():
    wb = xlsxwriter.Workbook(XLSX)
    title = wb.add_format({"bold": True, "font_name": FONT, "font_size": 14, "font_color": WHITE, "bg_color": NAVY, "valign": "vcenter"})
    header = wb.add_format({"bold": True, "font_name": FONT, "font_color": WHITE, "bg_color": BLUE, "text_wrap": True, "valign": "vcenter", "border": 1, "border_color": WHITE})
    cell = wb.add_format({"font_name": FONT, "font_size": 10, "text_wrap": True, "valign": "top", "border": 1, "border_color": "D9E1EA"})
    alt = wb.add_format({"font_name": FONT, "font_size": 10, "text_wrap": True, "valign": "top", "bg_color": LIGHT, "border": 1, "border_color": "D9E1EA"})
    note = wb.add_format({"font_name": FONT, "font_size": 10, "text_wrap": True, "valign": "top", "bg_color": PALE_TEAL})

    ws = wb.add_worksheet("口径说明")
    ws.set_column("A:A", 25)
    ws.set_column("B:B", 110)
    notes = [
        ("使用场景", "汇报PPT材料。包括会议指定FlagGems 10项、FlagGems-SGLang五项，以及早期已交付算子。"),
        ("性能口径", "HG-KDA成绩为KGS完整调用walltime相对等价PyTorch参考的几何平均。门槛为>0.8×；>1.0×才表示超过PyTorch。"),
        ("会议10项状态", "截至2026-10-26：7项默认条件达标、1项仅在官方确定性条件达标、2项只有CPU回退参考上的分数并完成补充NPU分解对照。不能说10项默认全达标。"),
        ("CPU回退", "fractional_max_pool3d与_conj_physical的17.683×、47.170×来自含CPU回退的PyTorch参考，不代表相对原生NPU kernel的加速；另列自建NPU分解对照10.066×、2.296×。"),
        ("方法归属", "KG/HG-KDA、CC/DeepSeek与Codex共同参与过历史优化、诊断或修复。结果表不表示HG-KDA独立自动取得全部成绩；具体贡献以来源材料为准。"),
        ("Excel记录", "从仓库内验收看板_结果表.xlsx直接读取华为Ascend910B原始单元格；保留‘精度失败/精度通过’等文字，不替换成推测性能数值。"),
        ("FlagGems版本", "会议FlagGems固定提交d6a8eec473517a3d68157b208eb9c057eb1d4c50，FlagTree 0.6.0+ascend.gitc286cba6，CANN 9.0。SGLang固定版本938612058a981487e439da514df7d1371cf31e68。"),
        ("selu_复测", "selu_在会议10项最终表中为0.837×；buffer-pointwise扩展批次另记录0.898×，来源批次和候选版本不同。分别保留原值，不跨批次平均。"),
        ("原始证据", "本工作簿只汇总，不改写原始实验、合同、源码或失败记录。末列给出项目内相对路径。"),
    ]
    ws.merge_range("A1:B1", "Ascend 910B 算子结果汇报口径", title)
    ws.set_row(0, 26)
    for i, (k, v) in enumerate(notes, start=2):
        ws.write(i - 1, 0, k, header)
        ws.write(i - 1, 1, v, note)
        ws.set_row(i - 1, 42)

    def add_table(name, subtitle, headers, rows, widths):
        sh = wb.add_worksheet(name)
        sh.merge_range(0, 0, 0, len(headers) - 1, subtitle, title)
        sh.set_row(0, 26)
        sh.write_row(1, 0, headers, header)
        sh.set_row(1, 34)
        for r, row in enumerate(rows, start=2):
            fmt = cell if r % 2 == 0 else alt
            sh.write_row(r, 0, row, fmt)
            sh.set_row(r, 72 if len(headers) > 6 else 62)
        for col, width in enumerate(widths):
            sh.set_column(col, col, width)
        sh.freeze_panes(2, 1)
        sh.autofilter(1, 0, len(rows) + 1, len(headers) - 1)
        sh.set_landscape()
        sh.fit_to_pages(1, 0)
        sh.repeat_rows(0, 1)
        return sh

    add_table(
        "FlagGems-会议10项",
        "会议指定10项｜按可交付性排序｜Excel与实测分列",
        ["算子", "Excel原记录", "FlagGems实测", "当前优化版／PyTorch", "正式与补充验证", "关键优化", "结论／限制", "证据来源"],
        meeting,
        [34, 23, 37, 43, 47, 55, 55, 72],
    )
    add_table(
        "FlagGems-SGLang",
        "FlagGems-SGLang选定五项｜无对应Excel列",
        ["算子", "FlagGems-SGLang实测", "当前优化版／PyTorch", "验证", "关键优化／口径", "状态", "证据来源"],
        sglang,
        [32, 40, 48, 55, 60, 48, 70],
    )
    add_table(
        "早期已交付",
        "早期已交付与流程验证算子｜历史结果单独保留",
        ["算子", "Excel原记录", "FlagGems实测", "当前优化版／PyTorch", "相对FlagGems", "完整验证", "关键优化", "证据来源"],
        historic,
        [33, 20, 36, 26, 25, 24, 58, 65],
    )
    wb.close()


def rgb(value):
    return RGBColor.from_string(value)


def style_run(run, size, color=NAVY, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def textbox(slide, x, y, w, h, text, size=16, color=NAVY, bold=False, align=None, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    style_run(run, size, color, bold)
    return shape


def add_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_heading(slide, section, title, subtitle=""):
    textbox(slide, 0.55, 0.25, 12.2, 0.28, section, 11, TEAL, True)
    textbox(slide, 0.55, 0.58, 12.2, 0.55, title, 25, NAVY, True)
    if subtitle:
        textbox(slide, 0.58, 1.16, 12.1, 0.45, subtitle, 12, GRAY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.58), Inches(1.62), Inches(12.15), Inches(0.025))
    line.fill.solid(); line.fill.fore_color.rgb = rgb("DCE5ED"); line.line.fill.background()


def add_footer(slide, text):
    textbox(slide, 0.58, 7.16, 12.0, 0.18, text, 8, GRAY)


def add_card(slide, x, y, w, h, title, body, accent=TEAL, body_size=14):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = rgb(LIGHT); box.line.color.rgb = rgb("DCE5ED")
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = rgb(accent); stripe.line.fill.background()
    textbox(slide, x + 0.18, y + 0.12, w - 0.3, 0.46, title, 16, accent, True)
    textbox(slide, x + 0.18, y + 0.64, w - 0.34, h - 0.75, body, body_size, NAVY)
    return box


def add_table(slide, x, y, w, h, headers, rows, widths=None, font_size=11, header_size=11, row_height=None):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        total = sum(widths)
        for i, frac in enumerate(widths):
            table.columns[i].width = Inches(w * frac / total)
    for c, text in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = text
        cell.fill.solid(); cell.fill.fore_color.rgb = rgb(BLUE)
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                style_run(run, header_size, WHITE, True)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(text)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(WHITE if r % 2 else LIGHT)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c not in (0,) else PP_ALIGN.CENTER
                for run in p.runs:
                    style_run(run, font_size, NAVY, False)
        if row_height:
            table.rows[r].height = Inches(row_height)
    table.rows[0].height = Inches(0.43)
    return shape


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Cover and key takeaway.
    slide = prs.slides.add_slide(blank); add_bg(slide, NAVY)
    textbox(slide, 0.85, 0.72, 11.7, 0.35, "ASCEND 910B · HG-KDA", 13, "8AD4CE", True)
    textbox(slide, 0.82, 1.55, 11.7, 1.1, "从硬件证据到可验证的算子优化", 31, WHITE, True)
    textbox(slide, 0.85, 2.78, 11.5, 0.75, "方法闭环与会议算子阶段结果", 23, "D7E4EF")
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(4.1), Inches(11.65), Inches(1.45))
    box.fill.solid(); box.fill.fore_color.rgb = rgb("203651"); box.line.color.rgb = rgb("41617F")
    textbox(slide, 1.1, 4.34, 11.1, 0.35, "7 项默认验收达标  ·  1 项确定性条件达标  ·  2 项补充 NPU 对照", 20, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 1.05, 4.91, 11.2, 0.36, "报告中的 0.8× 是项目验收线；超过 1× 才代表快于 PyTorch。", 14, "B8C9D9", False, PP_ALIGN.CENTER)
    textbox(slide, 0.87, 6.84, 11.5, 0.25, "数据截至 2026-10-28｜固定 FlagGems／FlagTree 与 KGS 完整调用评测", 10, "B8C9D9")

    # 2. Methods.
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_heading(slide, "第一部分｜方法", "HG-KDA 把硬件观察接到可验证的代码改动", "老师提供成对 profiling 工作流；我们的增量是把指标计算、结构化诊断、热点筛选和经验约束接入同一优化闭环。")
    cards = [
        ("① 成对采集", "同一输入下对照候选与 PyTorch；核对设备、输入和目标 kernel。", BLUE),
        ("② 指标结构化", "脚本汇总搬运字节／指令、Vector／Scalar／MTE时间、等待与冲突；计算比值和平均搬运粒度。", TEAL),
        ("③ 机制定位", "组织源码、TTIR／bcmlir、真机指标与参考 Roofline；筛选目标 kernel 的热点指令并映射源码位置。", BLUE),
        ("④ 形成假设", "Agent结合算子类型和经验知识提出方向；阈值只触发待验证假设，不直接证明瓶颈。", GOLD),
        ("⑤ 完整验证", "候选经过 KGS 正式合同、扩展语义和实际计时；复采证据检查预测，保留失败与来源。", TEAL),
    ]
    x0, gap, cw, cy, ch = 0.58, 0.16, 2.30, 2.05, 2.7
    for i, (title, body, accent) in enumerate(cards):
        add_card(slide, x0 + i * (cw + gap), cy, cw, ch, title, body, accent, 12)
        if i < len(cards) - 1:
            textbox(slide, x0 + i * (cw + gap) + cw, cy + 1.05, gap, 0.35, "→", 20, GRAY, True, PP_ALIGN.CENTER)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(5.15), Inches(11.85), Inches(1.08))
    box.fill.solid(); box.fill.fore_color.rgb = rgb(PALE_TEAL); box.line.color.rgb = rgb("B9DEDA")
    textbox(slide, 0.98, 5.31, 11.4, 0.34, "脚本做可重复的采集、聚合、筛选和校验；模型结合证据判断并选择修改；KGS 决定候选是否通过。", 15, NAVY, True)
    textbox(slide, 0.98, 5.72, 11.3, 0.32, "当前限制：经验 Roofline 不是物理峰值模型；0.7／0.5 等阈值是经验提示；热点映射不等于完整关键路径因果重建。", 11, GRAY)
    add_footer(slide, "依据：tools/profiling/hgkda 与老师提供的 msprof-op-compare.zip 工作流；当前为脚本＋知识＋Agent调用的组合，不是完全自动的独立诊断器。")

    # 3. One concrete method case.
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_heading(slide, "第一部分｜算子例子", "matmul_bias_activation：从对照指标走到复测", "这个例子用于展示三步闭环的可执行性；入选源码复现了已有成功方案，不是本轮新发现。")
    add_card(slide, 0.68, 1.92, 3.78, 1.52, "起点与结果", "起点 0.865× → 候选 0.884×（相对 PyTorch）\n同 15 个 case 的实际延迟改善 1.023×", TEAL, 14)
    add_card(slide, 4.76, 1.92, 3.78, 1.52, "Agent选择的动作", "共享相邻输出块的 B；尝试改变输出 tile 复用方向。候选使用通用模板生成。", BLUE, 14)
    add_card(slide, 8.84, 1.92, 3.78, 1.52, "验收", "起点与候选均 42/42 正式通过，另有 18＋24 项补充检查通过。", TEAL, 14)
    add_table(slide, 0.82, 3.82, 11.7, 2.13,
              ["复采后的观测", "候选／起点", "能支持的判断"],
              [
                  ["GM／L2／L1总字节", "约 1.00×", "总流量基本没变"],
                  ["UB读写；Vector读UB", "0.831×；0.727×", "片内访问和尾部处理组织发生变化"],
                  ["AIV MTE2指令；平均搬运粒度", "0.818×；1.217×", "搬运指令减少、单条平均搬运更多"],
                  ["静态bcmlir copy／alloc", "5处 → 4处", "静态分配点减少；不能单独证明因果"],
              ], [3.7, 2.3, 5.7], 11, 11, 0.40)
    textbox(slide, 0.88, 6.18, 11.5, 0.48, "解释边界：候选仍略慢于 PyTorch（0.884×）；模拟器只导出部分指令／源码映射后异常退出，不能称为完整关键路径分析。", 12, RED, True)
    add_footer(slide, "来源：operators/campaigns/agent-demo-20261027/三步演示-核验版.md。候选源码与历史成功源码同字节，本批证明的是新采集和工作流复现。")

    # 4. Seven standard acceptances.
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_heading(slide, "第二部分｜新达标算子", "会议指定 FlagGems 算子：7 项默认条件达标", "性能均为 KGS 完整调用 walltime 相对等价 PyTorch 参考的几何平均；Excel 历史值与固定 FlagGems 实测分别列出。")
    rows = [
        ["value_selecting_reduction_backward", "0.4935×", "0.759×", "1.635×", "96项完整通过；布局修复"],
        ["scatter_add", "精度失败", "0.258×", "1.900×", "18＋18项；分段扫描修复动态范围"],
        ["uniform", "精度通过", "0.141×", "1.325×", "18项；标准Philox逐位通过"],
        ["rad2deg_", "精度失败", "失败33/36", "1.117×", "36＋18项；原地语义与启动"],
        ["upsample_bilinear2d", "精度失败", "编译／布局失败", "0.915×", "24＋48＋36项；连续行分块"],
        ["log10_", "精度失败", "失败33/36", "0.875×", "36＋42＋12＋9项；stride与搬运"],
        ["selu_", "精度失败", "0.119×", "0.837×", "36＋42＋12＋9项；分块与精度安排"],
    ]
    add_table(slide, 0.62, 1.89, 12.1, 4.55,
              ["算子", "Excel原记录", "FlagGems实测", "优化版", "主要验证／改动"], rows,
              [4.0, 1.85, 2.4, 1.6, 4.2], 11, 11, 0.58)
    textbox(slide, 0.8, 6.48, 11.7, 0.54, "注意：0.8× 是项目门槛；0.837–0.915× 仍低于 PyTorch 的 1.0×。selu_另有扩展批次0.898×，按来源分列，不跨批次平均。", 11, NAVY, True)
    add_footer(slide, "来源：operators/会议10算子最终验收-20261026.md。完整门禁和固定版本见该报告及对应 campaign。")

    # 5. Conditional and fallback cases.
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_heading(slide, "第二部分｜验收边界", "剩余 3 项要按参考条件单列", "会议10项的正式结论是 7 项默认达标、1 项确定性条件达标、2 项完成补充 NPU 分解对照。")
    rows = [
        ["cumulative_trapezoid", "精度失败", "0.812×（确定性条件）", "确定性33/33；默认参考重复结果不自洽，不能列默认达标"],
        ["fractional_max_pool3d", "0.0266×；FlagGems扩展16/27", "17.683× CPU回退；10.066× NPU分解", "9/9正式＋27＋30补充；CPU回退分数不当作原生NPU性能"],
        ["_conj_physical", "精度失败；新接口原版失败", "47.170× CPU回退；2.296× NPU分解", "14/14正式＋34补充；只作补充分解对照"],
    ]
    add_table(slide, 0.68, 1.96, 12.0, 2.63,
              ["算子", "Excel／原版", "报告中的性能数值", "结论"], rows,
              [2.7, 2.4, 3.0, 4.9], 11, 11, 0.72)
    add_card(slide, 0.78, 4.95, 5.75, 1.2, "确定性条件", "累计积分需在隔离进程显式启用官方确定性算法；默认参考仍阻塞。", GOLD, 13)
    add_card(slide, 6.78, 4.95, 5.75, 1.2, "CPU回退参考", "池化与共轭的大倍数来自含CPU回退的参考。NPU分解由多项NPU原语组成，也不是官方单kernel对照。", RED, 13)
    add_footer(slide, "来源：operators/会议10算子最终验收-20261026.md；operators/campaigns/acceptance-three-20261026/精简汇报.md。")

    # 6. SGLang results in a separate table.
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_heading(slide, "第二部分｜FlagGems-SGLang", "SGLang 五项分开报告", "本组没有对应 Excel 列；只有 gelu_and_mul 建议继续做 HG-KDA 深入优化。")
    rows = [
        ["gelu_and_mul", "0.068×", "2.271×", "19＋18＋15＋9项通过", "对齐本机GELU公式；连续二维分块", "已交付；活跃优化对象"],
        ["moe_sum_reduce", "0.965×", "1.624×", "两次8项正式；36项扩展", "融合归约；修复非连续布局", "已交付；生成＋Codex修复"],
        ["per_token_quant_int8", "2.085×", "—", "原版19/19", "原版体检已超过PyTorch", "归档，不重复优化"],
        ["rotary_embedding", "2.130×", "—", "原版9/9", "原版体检已超过PyTorch", "归档，不重复优化"],
        ["silu_and_mul_masked", "28.947×¹", "—", "参考口径观察", "组合PyTorch参考含循环和.item()同步", "不作公平融合kernel结论"],
    ]
    add_table(slide, 0.65, 1.93, 12.0, 3.95,
              ["算子", "FlagGems-SGLang", "优化版／PyTorch", "评测", "主要优化／口径", "状态"], rows,
              [2.45, 2.0, 2.0, 2.55, 3.35, 2.25], 10, 10, 0.62)
    textbox(slide, 0.8, 6.11, 11.8, 0.48, "¹ silu_and_mul_masked 的 28.947× 是带循环与同步的组合参考，不作为公平融合 kernel 的加速结论。", 11, GRAY)
    add_footer(slide, "来源：docs/实验归属与汇报口径.md；operators/campaigns/night-20261021-v3/精简汇报.md。")

    prs.save(PPT)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_xlsx()
    build_ppt()
    print(PPT)
    print(XLSX)
