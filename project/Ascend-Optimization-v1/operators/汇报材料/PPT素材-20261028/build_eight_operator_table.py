"""从已核对汇报数据生成7项会议算子与1项SGLang算子的PPT表格。"""
from pathlib import Path
import ast
import json
import zipfile
import xml.etree.ElementTree as ET
import xlsxwriter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
STEM = '本周8算子-7项会议加1项SGLang-20261028'
TITLE = '会议指定算子：7/10 项达标＋SGLang 1 项达标'
NAVY, TEAL, PALE, ALT = '19375C', '087F8C', 'E4F3EF', 'F1F5F9'
HEAD = ['算子', 'Excel 原记录', 'FlagGems 实测', 'HG-KDA', '完整评测', '关键优化']
SGHEAD = ['算子', 'FlagGems-SGLang 实测', 'HG-KDA', '完整评测', '关键优化']
NOTE = '相对PyTorch的加速比（KGS完整调用计时）；>0.8×达标，>1×快于PyTorch。完整评测列为正式用例，补充检查亦已通过。'
tree = ast.parse((OUT/'build_ppt_materials.py').read_text(encoding='utf-8'))
data = {n.targets[0].id: ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.Assign) and isinstance(n.targets[0],ast.Name) and n.targets[0].id in ('meeting','sglang')}
COUNTS = ['96/96','18/18','18/18','36/36','24/24','36/36','36/36']
BASE = [0.759,'0.258×；补充检查失败','0.141×；补充检查失败','正确性失败','编译／布局失败','正确性失败',0.119]
ACTIONS = [
 '融合补零与梯度回填；修复非连续布局',
 '排序后按目标分段累加，避免跨目标数值干扰',
 '修复随机数整数运算；优化状态读取与启动',
 '修复原地计算语义；调整分块与启动',
 '连续读取相邻输入行，一次算两行；修复交界重复写',
 '修复高维寻址；减少中间类型转换，优化启动',
 '调整分块与精度安排，优化启动',
]
def number(s): return float(s.split('×')[0])
ROWS=[]
for i,r in enumerate(data['meeting'][:7]):
    excel=number(r[1]) if r[1].endswith('×') else r[1]
    ROWS.append([r[0],excel,BASE[i],number(r[3]),COUNTS[i],ACTIONS[i]])
SG=[['gelu_and_mul','0.068×；补充检查失败',number(data['sglang'][0][2]),'19/19','对齐本机GELU公式；连续二维分块，简化索引与启动']]
GROUPS=[('会议7项',HEAD,ROWS,3),('SGLang1项',SGHEAD,SG,2)]
def shown(v,excel=False): return (f'{v:.4f}×' if excel else f'{v:.3f}×') if isinstance(v,float) else str(v)

def workbook():
    path=OUT/(STEM+'.xlsx'); wb=xlsxwriter.Workbook(path)
    for name,headers,rows,hg in GROUPS:
        ws=wb.add_worksheet(name); ws.hide_gridlines(2); ws.set_column(0,0,43)
        for c in range(1,len(headers)-1): ws.set_column(c,c,23 if c!=hg else 16)
        ws.set_column(len(headers)-1,len(headers)-1,62)
        header=wb.add_format({'font_name':'Microsoft YaHei','bold':True,'font_color':'white','bg_color':NAVY,'valign':'vcenter','align':'center'})
        ws.write_row(0,0,headers,header); ws.set_row(0,32)
        for r,row in enumerate(rows,1):
            ws.set_row(r,58)
            for c,v in enumerate(row):
                opts={'font_name':'Microsoft YaHei','font_size':11,'text_wrap':True,'valign':'vcenter','align':'left' if c in (0,len(headers)-1) else 'center','bg_color':PALE if c==hg else ALT if r%2==0 else 'FFFFFF','font_color':TEAL if c==hg else NAVY,'bold':c==hg,'bottom':1,'bottom_color':'DCE5EE'}
                if isinstance(v,float): opts['num_format']='0.0000"×"' if name=='会议7项' and c==1 else '0.000"×"'
                ws.write(r,c,v,wb.add_format(opts))
        ws.set_landscape(); ws.fit_to_pages(1,1); ws.print_area(0,0,len(rows),len(headers)-1)
    wb.close(); return path

def ppt():
    prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    def text(x,y,w,h,value,size=16,bold=False):
        sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.word_wrap=True
        run=tf.paragraphs[0].add_run(); run.text=value; run.font.name='Microsoft YaHei'; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor.from_string(NAVY)
    text(.45,.18,15.1,.65,TITLE,29,True)
    text(.46,.88,15,.37,'FlagGems会议指定算子｜默认条件下完整评测通过，均超过0.8×验收线',14)
    def table(y,headers,rows,widths,hg,height):
        tb=slide.shapes.add_table(len(rows)+1,len(headers),Inches(.45),Inches(y),Inches(15.1),Inches(height)).table
        for c,w in enumerate(widths): tb.columns[c].width=Inches(w)
        for r,row in enumerate([headers]+rows):
            tb.rows[r].height=Inches(.48 if r==0 else (height-.48)/len(rows))
            for c,v in enumerate(row):
                cell=tb.cell(r,c); cell.text=shown(v,r>0 and len(headers)==6 and c==1)
                cell.margin_left=cell.margin_right=Inches(.09); cell.margin_top=cell.margin_bottom=Inches(.03); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor.from_string((TEAL if c==hg else NAVY) if r==0 else PALE if c==hg else ALT if r%2==0 else 'FFFFFF')
                for p in cell.text_frame.paragraphs:
                    p.alignment=PP_ALIGN.LEFT if c in (0,len(headers)-1) else PP_ALIGN.CENTER
                    for run in p.runs:
                        run.font.name='Microsoft YaHei'; run.font.size=Pt(13 if r and c==0 else 14); run.font.bold=r==0 or c==hg; run.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else TEAL if c==hg else NAVY)
    table(1.34,HEAD,ROWS,[3.7,1.7,2.1,1.3,1.1,5.2],3,5.28)
    text(.46,6.72,15,.36,'FlagGems-SGLang｜独立列示',14,True)
    table(7.15,SGHEAD,SG,[3.7,3.8,1.3,1.1,5.2],2,1.05)
    text(.46,8.37,15,.4,NOTE,10)
    path=OUT/(STEM+'.pptx'); prs.save(path); return path

def png():
    im=Image.new('RGB',(2560,1500),'white'); d=ImageDraw.Draw(im)
    def font(size,bold=False): return ImageFont.truetype('C:/Windows/Fonts/'+('msyhbd.ttc' if bold else 'msyh.ttc'),size)
    d.text((64,35),TITLE,font=font(57,True),fill='#'+NAVY)
    d.text((65,124),'FlagGems会议指定算子｜默认条件下完整评测通过，均超过0.8×验收线',font=font(29),fill='#'+NAVY)
    def wrap(value,f,w):
        lines=['']
        for ch in value:
            if ch=='\n': lines.append(''); continue
            if d.textlength(lines[-1]+ch,font=f)>w and lines[-1]: lines.append(ch)
            else: lines[-1]+=ch
        return '\n'.join(lines)
    def table(y,headers,rows,widths,hg,rh):
        for r,row in enumerate([headers]+rows):
            x=64; h=88 if r==0 else rh
            for c,v in enumerate(row):
                fill=(TEAL if c==hg else NAVY) if r==0 else PALE if c==hg else ALT if r%2==0 else 'FFFFFF'
                d.rectangle((x,y,x+widths[c],y+h),fill='#'+fill)
                if r: d.line((x,y+h,x+widths[c],y+h),fill='#DCE5EE',width=2)
                f=font(31 if r==0 else 29,r==0 or c==hg); value=shown(v,r>0 and len(headers)==6 and c==1)
                if r and c==0 and value=='value_selecting_reduction_backward': value='value_selecting_\nreduction_backward'
                value=wrap(value,f,widths[c]-32); box=d.multiline_textbbox((0,0),value,font=f,spacing=7); th=box[3]-box[1]
                tx=x+16 if c in (0,len(headers)-1) else x+(widths[c]-(box[2]-box[0]))/2
                d.multiline_text((tx,y+(h-th)/2-box[1]),value,font=f,spacing=7,fill='#'+('FFFFFF' if r==0 else TEAL if c==hg else NAVY))
                x+=widths[c]
            y+=h
        return y
    end=table(190,HEAD,ROWS,[540,280,340,210,180,882],3,120)
    d.text((65,end+28),'FlagGems-SGLang｜独立列示',font=font(30,True),fill='#'+NAVY)
    table(end+86,SGHEAD,SG,[540,620,210,180,882],2,115)
    d.text((65,1430),NOTE,font=font(24),fill='#526174')
    path=OUT/(STEM+'.png'); im.save(path); return path

def markdown():
    lines=['# '+TITLE,'',NOTE,'']
    for name,headers,rows,hg in GROUPS:
        lines+=['## '+name,'','| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']
        for row in rows: lines.append('| '+' | '.join(shown(v,name=='会议7项' and c==1) for c,v in enumerate(row))+' |')
        lines+=['']
    lines+=['会议7项包含原记录正确性失败的算子，故标题不写“原加速比低于0.8×的8算子”。这是当前累计交付表，不表示8项均为本周新增。','','FlagGems固定提交：d6a8eec473517a3d68157b208eb9c057eb1d4c50；FlagGems-SGLang：938612058a981487e439da514df7d1371cf31e68；FlagTree：0.6.0+ascend.gitc286cba6。','','HG-KDA列按项目口径汇总KG/HG-KDA、CC/DeepSeek与Codex参与的有效交付，不表示全部由独立Agent自动取得。','','数据来源：build_ppt_materials.py；operators/会议10算子最终验收-20261026.md；operators/夜间连续优化交付-20261023.md；operators/两算子收尾交付-20261024.md。']
    path=OUT/(STEM+'.md'); path.write_text('\n'.join(lines)+'\n',encoding='utf-8'); return path

if __name__=='__main__':
    paths=[workbook(),ppt(),png(),markdown()]
    loaded=Presentation(paths[1]); tables=[s.table for s in loaded.slides[0].shapes if s.has_table]
    assert len(tables)==2 and len(tables[0].rows)==8 and len(tables[1].rows)==2
    assert [tables[0].cell(r,0).text for r in range(1,8)]==[r[0] for r in ROWS]
    assert tables[1].cell(1,0).text=='gelu_and_mul'
    with zipfile.ZipFile(paths[0]) as z:
        sheets=[n for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')]
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        assert len(sheets)==2
        assert [len(ET.fromstring(z.read(s)).findall('.//s:sheetData/s:row',ns)) for s in sheets]==[8,2]
        assert float(ET.fromstring(z.read(sheets[0])).find('.//s:c[@r="D2"]/s:v',ns).text)==1.635
    print(json.dumps({'files':[str(p) for p in paths],'operators':8,'verified':'PPT与Excel两表7+1、算子顺序、数值类型与首项加速比核对通过'},ensure_ascii=False))
