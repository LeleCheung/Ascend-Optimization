from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import xlsxwriter  # noqa: F401 - confirms the generator dependency is importable.
from pptx import Presentation


root = Path(__file__).resolve().parent
ppt = root / "Ascend910B_HG-KDA汇报-20261028.pptx"
xlsx = root / "算子结果总表-20261028.xlsx"

with zipfile.ZipFile(ppt) as archive:
    assert archive.testzip() is None, "PPTX ZIP integrity failure"
with zipfile.ZipFile(xlsx) as archive:
    assert archive.testzip() is None, "XLSX ZIP integrity failure"
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    sheets = [node.attrib["name"] for node in workbook.findall("m:sheets/m:sheet", ns)]
    assert sheets == ["口径说明", "FlagGems-会议10项", "FlagGems-SGLang", "早期已交付"], sheets
    row_counts = []
    for number in range(1, 5):
        sheet_xml = ET.fromstring(archive.read(f"xl/worksheets/sheet{number}.xml"))
        row_counts.append(len(sheet_xml.findall("m:sheetData/m:row", ns)))
    assert row_counts == [10, 12, 7, 11], row_counts

prs = Presentation(ppt)
assert len(prs.slides) == 6, len(prs.slides)
all_text = []
for slide_number, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (slide_number, shape.name, "negative position")
        assert shape.left + shape.width <= prs.slide_width + 2, (slide_number, shape.name, "horizontal overflow")
        assert shape.top + shape.height <= prs.slide_height + 2, (slide_number, shape.name, "vertical overflow")
        if getattr(shape, "has_text_frame", False):
            all_text.append(shape.text)
        if getattr(shape, "has_table", False):
            all_text.extend(cell.text for row in shape.table.rows for cell in row.cells)

text = "\n".join(all_text)
for phrase in ["7 项默认条件达标", "matmul_bias_activation", "1.023×", "0.884×", "value_selecting_reduction_backward", "cumulative_trapezoid", "CPU回退", "silu_and_mul_masked"]:
    assert phrase in text, phrase

print("PPTX: 6 slides, all shapes within slide bounds, expected content present")
print("XLSX: valid package, sheets:", ", ".join(sheets), "rows:", row_counts)
